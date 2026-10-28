"""Build PMR SI abstract + late-entry email DOCX from the canonical .txt files."""
import os
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))


def build(txt_name, docx_name, title):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = None
    with open(os.path.join(HERE, txt_name), encoding="utf-8") as f:
        text = f.read()
    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        if para.startswith("Subject:") or para.startswith("Proposed title:") or para.startswith("Author:"):
            p = doc.add_paragraph()
            run = p.add_run(para)
            run.bold = True
        else:
            doc.add_paragraph(para)
    out = os.path.join(HERE, docx_name)
    doc.save(out)
    print("wrote", out)


build(
    "PMR_SI_ABSTRACT_PROPOSAL_KYOTO_SHIGA_FINAL.txt",
    "PMR_SI_ABSTRACT_PROPOSAL_KYOTO_SHIGA_FINAL.docx",
    "Abstract proposal — PMR SI Place-Based Strategy Design for Value Creation",
)
build(
    "PMR_SI_LATE_ENTRY_INQUIRY_KYOTO_SHIGA_FINAL.txt",
    "PMR_SI_LATE_ENTRY_INQUIRY_KYOTO_SHIGA_FINAL.docx",
    "Late-entry inquiry — PMR SI abstract proposal",
)
