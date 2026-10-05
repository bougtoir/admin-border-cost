"""Build highlights.docx from the Highlights section of manuscript_text.md."""
from pathlib import Path

from docx import Document

from .docx_rich import add_rich_text, normalize_styles

ROOT = Path(__file__).resolve().parents[2]
MAN = ROOT / "manuscript"


def main():
    txt = (MAN / "manuscript_text.md").read_text()
    sec = txt.split("## Highlights")[1].split("## Abstract")[0]
    bullets = [l.strip()[2:] for l in sec.splitlines() if l.strip().startswith("- ")]
    doc = Document()
    doc.add_heading("Highlights", 0)
    for b in bullets:
        p = doc.add_paragraph()
        add_rich_text(p, b)
    normalize_styles(doc)
    doc.save(MAN / "highlights.docx")
    for b in bullets:
        assert len(b) <= 85, f"highlight over 85 chars: {b}"
    print(f"highlights.docx written ({len(bullets)} bullets)")


if __name__ == "__main__":
    main()
