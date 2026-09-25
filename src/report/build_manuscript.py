"""Build manuscript DOCX from manuscript_text.md template + computed values.

Every number in prose is substituted from manuscript_values.csv ({{key}}),
so no value is hard-coded. Figures are NOT embedded (journal submission:
figures as separate files) — figure/table captions are listed instead.
"""
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
MAN = ROOT / "manuscript"
VALS = pd.read_csv(MAN / "manuscript_values.csv").set_index("value_id").display_value \
    if (MAN / "manuscript_values.csv").exists() else pd.Series(dtype=object)

FMT = {
    # keys formatted as numbers-with-commas etc.
}


def val(key):
    if key not in VALS.index or pd.isna(VALS[key]):
        return f"[MISSING {key}]"
    v = VALS[key]
    if isinstance(v, str):
        try:
            v = float(v)
        except ValueError:
            return v
    if isinstance(v, float):
        if v == int(v) or abs(v) >= 1000:
            return f"{v:,.0f}"
        return f"{v:.3f}"
    return str(v)


def fill(text):
    return re.sub(r"\{\{(\w+)\}\}", lambda m: val(m.group(1)), text)


def main():
    tpl = (MAN / "manuscript_text.md").read_text()
    body = fill(tpl)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for line in body.splitlines():
        s = line.rstrip()
        if s.startswith("# "):
            doc.add_heading(s[2:], 0)
        elif s.startswith("## "):
            doc.add_heading(s[3:], 1)
        elif s.startswith("### "):
            doc.add_heading(s[4:], 2)
        elif s.strip():
            doc.add_paragraph(s)
    doc.save(MAN / "manuscript.docx")
    miss = re.findall(r"\[MISSING \w+\]", body)
    print("manuscript.docx written; missing values:", sorted(set(miss)))


if __name__ == "__main__":
    main()
