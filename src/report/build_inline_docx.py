"""Build manuscript_inline.docx: manuscript body with figures and tables
embedded inline (for internal review; the journal version keeps figures
as separate files). Values filled from manuscript_values.csv via the same
placeholder mechanism as build_manuscript.
"""
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
MAN = ROOT / "manuscript"
FIG = ROOT / "outputs" / "figures"
TAB = ROOT / "outputs" / "tables"

VALS = pd.read_csv(MAN / "manuscript_values.csv").set_index("value_id").display_value


def val(key):
    if key not in VALS.index or pd.isna(VALS[key]):
        return f"[MISSING {key}]"
    v = VALS[key]
    if isinstance(v, str):
        try:
            v = float(v)
        except ValueError:
            return v
    if isinstance(v, float):
        if v == int(v):
            return f"{v:,.0f}"
        return f"{v:,.0f}" if abs(v) >= 1000 else f"{v:.3f}"
    return str(v)


def fill(text):
    return re.sub(r"\{\{(\w+)\}\}", lambda m: val(m.group(1)), text)


FIG_FILES = [
    ("Figure 1. Study region: Kyoto-fu and Shiga-ken, census small areas (JGD2011).", "F1_study_area.png"),
    ("Figure 2. Current statutory House-of-Representatives districts.", "F2_current_map.png"),
    ("Figure 3. Best-found maps with the border constraint ON vs OFF.", "F3_optimal_on_off.png"),
    ("Figure 4. Cross-prefecture commuting/schooling OD network.", "F4_flow_network.png"),
    ("Figure 5. Objective-component comparison, ON vs OFF.", "F5_metric_comparison.png"),
    ("Figure 6. Placebo-border cost distribution vs the real border.", "F6_placebo.png"),
    ("Figure 7. Future robustness under IPSS projections to 2050.", "F7_future_robustness.png"),
]

TABLES = [
    ("Table 1. Data sources.", "T1_data_sources.csv"),
    ("Table 2. Current-map metrics.", "T2_current_metrics.csv"),
    ("Table 3. Model comparison ON vs OFF.", "T3_on_off.csv"),
    ("Table 4. tau sensitivity.", "T4_tau_sensitivity.csv"),
    ("Table 5. Placebo summary.", "T5_placebo.csv"),
    ("Table 6. Future robustness.", "T6_future.csv"),
]


def add_table(doc, csv_path):
    df = pd.read_csv(csv_path)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    for j, c in enumerate(df.columns):
        t.rows[0].cells[j].text = str(c)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for j, x in enumerate(row):
            cells[j].text = str(x)


def main():
    tpl = (MAN / "manuscript_text.md").read_text()
    body = fill(tpl)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    fig_inserted = set()
    tab_inserted = set()
    skip_lists = False
    for line in body.splitlines():
        s = line.rstrip()
        if s.startswith("## Figures") or s.startswith("## Tables"):
            skip_lists = True   # replace the caption lists with embedded content
            continue
        if s.startswith("## "):
            skip_lists = False
        if skip_lists and s.strip().startswith("- "):
            continue
        if s.startswith("# "):
            doc.add_heading(s[2:], 0)
        elif s.startswith("## "):
            doc.add_heading(s[3:], 1)
        elif s.startswith("### "):
            doc.add_heading(s[4:], 2)
        elif s.strip():
            doc.add_paragraph(s)
        # embed each figure/table right after the paragraph that first cites it
        for cap, fn in FIG_FILES:
            tag = fn.split("_")[0]          # "F1".."F7"
            if tag not in fig_inserted and (f"Figure {tag}" in s or f"({tag}" in s) and (FIG / fn).exists():
                p = doc.add_paragraph(cap)
                p.runs[0].bold = True
                doc.add_picture(str(FIG / fn), width=Inches(6.0))
                fig_inserted.add(tag)
        for cap, fn in TABLES:
            tag = fn.split("_")[0]          # "T1".."T6"
            if tag not in tab_inserted and f"{tag}" in s and (TAB / fn).exists():
                p = doc.add_paragraph(cap)
                p.runs[0].bold = True
                add_table(doc, TAB / fn)
                tab_inserted.add(tag)
    # anything uncited still gets appended
    doc.add_page_break()
    doc.add_heading("Remaining figures and tables", 1)
    for cap, fn in FIG_FILES:
        tag = fn.split("_")[0]
        if tag not in fig_inserted and (FIG / fn).exists():
            p = doc.add_paragraph(cap)
            p.runs[0].bold = True
            doc.add_picture(str(FIG / fn), width=Inches(6.0))
    for cap, fn in TABLES:
        tag = fn.split("_")[0]
        if tag not in tab_inserted and (TAB / fn).exists():
            p = doc.add_paragraph(cap)
            p.runs[0].bold = True
            add_table(doc, TAB / fn)
    doc.save(MAN / "manuscript_inline.docx")
    miss = re.findall(r"\[MISSING \w+\]", body)
    print("manuscript_inline.docx written; missing:", sorted(set(miss)))


if __name__ == "__main__":
    main()
