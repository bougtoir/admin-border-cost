"""Rich text -> docx paragraph: **bold**, *italic*, and math tokens -> OMML.

Math token rules are an explicit ordered list (no greedy detection): only
declared patterns become Word equations; every other asterisk stays as the
markdown italic marker it already is.
"""
import re

from docx_equation import mathml_to_omml
from latex2mathml.converter import convert as latex_to_mathml

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _omml(p, latex):
    p._element.append(mathml_to_omml(latex_to_mathml(latex)))


# ordered (pattern, latex-builder); matched spans never reach the italic pass
MATH_RULES = [
    # case-selection flow notation (longest first)
    (re.compile(r"R_\{i→j\} = F_\{i→j\}/T_i"),
     lambda m: r"R_{i \to j} = \frac{F_{i \to j}}{T_i}"),
    (re.compile(r"S_\{ij\} = F_\{ij\}\^\{bi\}/\(P_i\+P_j\)"),
     lambda m: r"S_{ij} = \frac{F_{ij}^{bi}}{P_i + P_j}"),
    (re.compile(r"F_\{i→j\}"), lambda m: r"F_{i \to j}"),
    (re.compile(r"F_\{j→i\}"), lambda m: r"F_{j \to i}"),
    (re.compile(r"F_\{ij\}\^\{bi\}"), lambda m: r"F_{ij}^{bi}"),
    (re.compile(r"R_\{i→j\}"), lambda m: r"R_{i \to j}"),
    (re.compile(r"R_\{j→i\}"), lambda m: r"R_{j \to i}"),
    (re.compile(r"S_\{ij\}"), lambda m: r"S_{ij}"),
    (re.compile(r"R_ext\b"), lambda m: r"R_{\mathrm{ext}}"),
    (re.compile(r"S_2\b"), lambda m: r"S_2"),
    (re.compile(r"T_i\b"), lambda m: r"T_i"),
    (re.compile(r"P_i\b"), lambda m: r"P_i"),
    (re.compile(r"P_j\b"), lambda m: r"P_j"),
    # legacy underscore forms, if any survive in prose
    (re.compile(r"F_ij\b"), lambda m: r"F_{i \to j}"),
    (re.compile(r"F_ji\b"), lambda m: r"F_{j \to i}"),
    (re.compile(r"F_bi\b"), lambda m: r"F_{ij}^{bi}"),
    (re.compile(r"R_ij\b"), lambda m: r"R_{i \to j}"),
    # estimand / feasible-set notation
    (re.compile(r"C_s\(B\) = L_s\*\(B\) − L_s\*\(0\) ≥ 0"),
     lambda m: r"C_s(B) = L_s^{*}(B) - L_s^{*}(0) \ge 0"),
    (re.compile(r"D_s\* = argmin_D L_s\(D\)"),
     lambda m: r"D_s^{*} = \arg\min_{D} L_s(D)"),
    (re.compile(r"s ∈ \{1, …, S\}"),
     lambda m: r"s \in \{1, \ldots, S\}"),
    (re.compile(r"F_B ⊂ F_0"),
     lambda m: r"F_B \subset F_0"),
    (re.compile(r"L_s\*\(B\)"), lambda m: r"L_s^{*}(B)"),
    (re.compile(r"L_s\*\(0\)"), lambda m: r"L_s^{*}(0)"),
    (re.compile(r"C_s\(B\)"), lambda m: r"C_s(B)"),
    (re.compile(r"L_s\(D\)"), lambda m: r"L_s(D)"),
    (re.compile(r"L_s\b"), lambda m: r"L_s"),
    (re.compile(r"D_s\*"), lambda m: r"D_s^{*}"),
    (re.compile(r"C\(B\)"), lambda m: r"C(B)"),
    (re.compile(r"F_B"), lambda m: r"F_B"),
    (re.compile(r"F_0"), lambda m: r"F_0"),
    (re.compile(r"Cost_border = L\*\(Boundary ON\) − L\*\(Boundary OFF\)"),
     lambda m: r"C_{\mathrm{border}} = L^{*}\!\left(\mathrm{Boundary\ ON}\right) - L^{*}\!\left(\mathrm{Boundary\ OFF}\right)"),
    (re.compile(r"Cost_border = L\*\(ON\) − L\*\(OFF\)"),
     lambda m: r"C_{\mathrm{border}} = L^{*}\!\left(\mathrm{ON}\right) - L^{*}\!\left(\mathrm{OFF}\right)"),
    (re.compile(r"L\*_OFF ≤ L\*_ON"),
     lambda m: r"L^{*}_{\mathrm{OFF}} \le L^{*}_{\mathrm{ON}}"),
    (re.compile(r"L\*_(OFF|ON)"),
     lambda m: rf"L^{{*}}_{{\mathrm{{{m.group(1)}}}}}"),
    (re.compile(r"Cost_border\(([AB])\)"),
     lambda m: rf"C_{{\mathrm{{border}}}}({m.group(1)})"),
    (re.compile(r"Cost_border"),
     lambda m: r"C_{\mathrm{border}}"),
    (re.compile(r"±τ"),
     lambda m: r"\pm\tau"),
    (re.compile(r"±(\d+(?:\.\d+)?)%(?!\w)"),
     lambda m: rf"\pm {m.group(1)}\%"),
    (re.compile(r"±(\d+(?:\.\d+)?)\s*(km|m)\b"),
     lambda m: rf"\pm {m.group(1)}\,\mathrm{{{m.group(2)}}}"),
    (re.compile(r"τ"),
     lambda m: r"\tau"),
    (re.compile(r"max\|dev\|"),
     lambda m: r"|\mathrm{dev}|_{\max}"),
    (re.compile(r"max/min"),
     lambda m: r"\max/\min"),
]

_MATH_SPAN_RE = re.compile("|".join(f"(?:{r.pattern})" for r, _ in MATH_RULES))
_RULE_RES = [r for r, _ in MATH_RULES]
_RULE_FNS = [f for _, f in MATH_RULES]


def _latex_for(span):
    for r, fn in zip(_RULE_RES, _RULE_FNS):
        m = r.fullmatch(span)
        if m:
            return fn(m)
    return None


def _add_italic_runs(p, text, bold=False):
    last = 0
    for m in re.finditer(r"\*([^*]+)\*", text):
        if m.start() > last:
            run = p.add_run(text[last:m.start()])
            run.bold = bold
        run = p.add_run(m.group(1))
        run.italic = True
        run.bold = bold
        last = m.end()
    if last < len(text):
        run = p.add_run(text[last:])
        run.bold = bold


def _add_segment(p, text, bold=False):
    last = 0
    for m in _MATH_SPAN_RE.finditer(text):
        latex = _latex_for(m.group(0))
        if latex is None:
            continue
        if m.start() > last:
            _add_italic_runs(p, text[last:m.start()], bold=bold)
        _omml(p, latex)
        last = m.end()
    if last < len(text):
        _add_italic_runs(p, text[last:], bold=bold)


def add_rich_text(p, text):
    """Render markdown inline (**bold**, *italic*) + declared math to OMML."""
    last = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > last:
            _add_segment(p, text[last:m.start()])
        _add_segment(p, m.group(1), bold=True)
        last = m.end()
    if last < len(text):
        _add_segment(p, text[last:])


from docx.shared import Pt, RGBColor

BLACK = RGBColor(0, 0, 0)
STYLE_SPECS = {
    "Normal": (12, None),
    "Title": (24, True),
    "Subtitle": (12, False),
    "Heading 1": (16, True),
    "Heading 2": (14, True),
    "Heading 3": (12, True),
    "Caption": (10, None),
    "Hyperlink": (12, None),
    "FollowedHyperlink": (12, None),
}


def normalize_styles(doc):
    """SEPS typography: Times New Roman everywhere, black text, 12 pt body,
    10 pt captions. Corrects the underlying style definitions (removes the
    default Word theme blue from headings/captions/hyperlinks)."""
    for name, (size, bold) in STYLE_SPECS.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        f = st.font
        f.name = "Times New Roman"
        # East Asian / complex-script font too, via rPr rFonts
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for a in list(rfonts.attrib):
            if a.endswith("Theme"):
                del rfonts.attrib[a]
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}",
                       "Times New Roman")
        f.size = Pt(size)
        if bold is not None:
            f.bold = bold
        f.color.rgb = BLACK
        # strip theme color attribute if present
        rpr2 = st.element.get_or_add_rPr()
        color_el = rpr2.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color")
        if color_el is not None:
            for a in list(color_el.attrib):
                if a.endswith("themeColor") or a.endswith("themeShade") or a.endswith("themeTint"):
                    del color_el.attrib[a]
