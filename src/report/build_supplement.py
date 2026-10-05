"""supplement.md -> supplement.docx (same renderer as the manuscript)."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .docx_rich import add_rich_text, normalize_styles

ROOT = Path(__file__).resolve().parents[2]
MAN = ROOT / "manuscript"


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for line in (MAN / "supplement.md").read_text().splitlines():
        s = line.rstrip()
        if s.startswith("# "):
            doc.add_heading(s[2:], 0)
        elif s.startswith("## "):
            doc.add_heading(s[3:], 1)
        elif s.strip():
            p = doc.add_paragraph()
            add_rich_text(p, s)
    normalize_styles(doc)
    doc.save(MAN / "supplement.docx")
    print("supplement.docx written")


if __name__ == "__main__":
    main()
